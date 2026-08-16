"""BOM export to .xlsx (SPEC.md §7)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from harnessboard.bom import BomRow, CableRow

_HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
_BOLD = Font(bold=True)


def export_bom(
    path: str | Path,
    parts: list[BomRow],
    cables: list[CableRow],
    harness_info: dict | None = None,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"

    row = 1
    ws.cell(row, 1, "HarnessBoard — Bill of Materials").font = Font(bold=True, size=14)
    row += 1
    info = harness_info or {}
    for label, key in (
        ("Customer / project", "customer"),
        ("Harness number", "number"),
        ("Harness type", "type"),
    ):
        if info.get(key):
            ws.cell(row, 1, label).font = _BOLD
            ws.cell(row, 2, str(info[key]))
            row += 1
    ws.cell(row, 1, "Date").font = _BOLD
    ws.cell(row, 2, date.today().isoformat())
    row += 2

    ws.cell(row, 1, "PARTS").font = _BOLD
    row += 1
    for col, title in enumerate(["Model name", "Category", "Quantity", "Z depth (mm)"], 1):
        cell = ws.cell(row, col, title)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
    row += 1
    for r in parts:
        ws.cell(row, 1, r.name)
        ws.cell(row, 2, r.category)
        ws.cell(row, 3, r.quantity)
        ws.cell(row, 4, r.z_depth if r.z_depth is not None else "")
        row += 1
    row += 1

    ws.cell(row, 1, "CABLES").font = _BOLD
    row += 1
    for col, title in enumerate(
        ["Designation", "Diameter (mm)", "Z standard (mm)"], 1
    ):
        cell = ws.cell(row, col, title)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
    row += 1
    for c in cables:
        ws.cell(row, 1, c.designation)
        ws.cell(row, 2, c.diameter)
        ws.cell(row, 3, c.z_standard if c.z_standard is not None else "")
        row += 1

    for col, width in enumerate((32, 22, 12, 16), 1):
        ws.column_dimensions[chr(ord("A") + col - 1)].width = width
    for line in ws.iter_rows():
        for cell in line:
            cell.alignment = Alignment(vertical="center")

    wb.save(str(path))
